"""投标附录模板：占位符替换、默认模板、打包下载。"""

from __future__ import annotations

import io
import json
import os
import re
import shutil
import tempfile
import uuid
import zipfile
from datetime import datetime
from typing import Any

import docx
from docx.shared import Pt

from ..utils.localdb_paths import localdb_company_root, sanitize_company_id

APPENDIX_DIR = "appendix"
PROFILE_FILE = "company_profile.json"

PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")

DEFAULT_TEMPLATES: dict[str, str] = {
    "投标函.docx": "投标函",
    "授权委托书.docx": "授权委托书",
    "承诺书.docx": "承诺书",
}


def appendix_root(company_id: str) -> str:
    return os.path.join(localdb_company_root(company_id), APPENDIX_DIR)


def save_company_profile(company_id: str, profile: dict[str, Any]) -> None:
    root = localdb_company_root(company_id)
    os.makedirs(root, exist_ok=True)
    path = os.path.join(root, PROFILE_FILE)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(profile, f, ensure_ascii=False, indent=2)


def load_company_profile(company_id: str) -> dict[str, Any]:
    path = os.path.join(localdb_company_root(company_id), PROFILE_FILE)
    if not os.path.isfile(path):
        return {}
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def build_placeholder_map(
    company_id: str,
    project_name: str | None = None,
    extra: dict[str, str] | None = None,
) -> dict[str, str]:
    p = load_company_profile(company_id)
    today = datetime.now().strftime("%Y年%m月%d日")
    mapping = {
        "company_name": str(p.get("company_name") or ""),
        "unified_code": str(p.get("unified_code") or ""),
        "established_date": str(p.get("established_date") or ""),
        "legal_representative": str(p.get("legal_representative") or ""),
        "registered_capital": str(p.get("registered_capital") or ""),
        "address": str(p.get("address") or ""),
        "enterprise_scale": str(p.get("enterprise_scale") or ""),
        "business_nature": str(p.get("business_nature") or ""),
        "project_name": (project_name or "").strip(),
        "date": today,
        "today": today,
    }
    if extra:
        mapping.update({k: str(v) for k, v in extra.items()})
    return mapping


def _replace_in_paragraph(paragraph, mapping: dict[str, str]) -> None:
    text = paragraph.text
    if not text or "{{" not in text:
        return

    def repl(match: re.Match) -> str:
        key = match.group(1)
        return mapping.get(key, match.group(0))

    new_text = PLACEHOLDER_PATTERN.sub(repl, text)
    if new_text == text:
        return
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = new_text
    else:
        paragraph.add_run(new_text)


def fill_docx_template(input_path: str, output_path: str, mapping: dict[str, str]) -> None:
    doc = docx.Document(input_path)
    for para in doc.paragraphs:
        _replace_in_paragraph(para, mapping)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_in_paragraph(para, mapping)
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    doc.save(output_path)


def _create_default_template(path: str, title: str) -> None:
    doc = docx.Document()
    h = doc.add_paragraph()
    r = h.add_run(title)
    r.bold = True
    r.font.size = Pt(16)
    h.alignment = 1
    doc.add_paragraph()
    lines = [
        "致：{{project_name}} 招标人",
        "",
        "投标人（盖章）：{{company_name}}",
        "统一社会信用代码：{{unified_code}}",
        "法定代表人（签字）：{{legal_representative}}",
        "地址：{{address}}",
        "日期：{{date}}",
    ]
    for line in lines:
        doc.add_paragraph(line)
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    doc.save(path)


def ensure_default_templates(company_id: str) -> None:
    root = appendix_root(company_id)
    os.makedirs(root, exist_ok=True)
    for filename, title in DEFAULT_TEMPLATES.items():
        full = os.path.join(root, filename)
        if not os.path.isfile(full):
            _create_default_template(full, title)


def list_appendix_templates(company_id: str) -> list[dict[str, Any]]:
    ensure_default_templates(company_id)
    root = appendix_root(company_id)
    out: list[dict[str, Any]] = []
    for fn in sorted(os.listdir(root)):
        if not fn.lower().endswith(".docx"):
            continue
        full = os.path.join(root, fn)
        if not os.path.isfile(full):
            continue
        stat = os.stat(full)
        out.append(
            {
                "name": fn,
                "size_bytes": int(stat.st_size),
                "mtime": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
            }
        )
    return out


def build_filled_appendix_zip(
    company_id: str,
    project_name: str | None = None,
    template_names: list[str] | None = None,
) -> io.BytesIO:
    cid = sanitize_company_id(company_id)
    ensure_default_templates(cid)
    root = appendix_root(cid)
    mapping = build_placeholder_map(cid, project_name=project_name)

    work = tempfile.mkdtemp(prefix="appendix_fill_")
    try:
        names = template_names or [
            fn for fn in os.listdir(root) if fn.lower().endswith(".docx")
        ]
        if not names:
            raise ValueError("没有可用的附录模板")

        for fn in names:
            src = os.path.join(root, fn)
            if not os.path.isfile(src):
                continue
            dst = os.path.join(work, fn)
            fill_docx_template(src, dst, mapping)

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for fn in os.listdir(work):
                full = os.path.join(work, fn)
                if os.path.isfile(full):
                    zf.write(full, arcname=fn)
        buf.seek(0)
        return buf
    finally:
        shutil.rmtree(work, ignore_errors=True)


def append_docx_files_to_document(main_doc: docx.Document, appendix_paths: list[str]) -> None:
    """将若干 docx 正文追加到主文档末尾（分页）。"""
    for path in appendix_paths:
        if not os.path.isfile(path):
            continue
        main_doc.add_page_break()
        src = docx.Document(path)
        for element in src.element.body:
            if element.tag.endswith("sectPr"):
                continue
            main_doc.element.body.append(element)
